import hashlib
import io
from pathlib import PurePosixPath
from zipfile import ZipFile

from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from starlette.concurrency import run_in_threadpool

from app.config import Settings
from app.errors import AppError
from app.models import Chunk, Course, Document


class DocumentParser:
    """Extract text only. No OCR or embedded file execution."""

    def parse(self, filename: str, data: bytes, max_chars: int) -> list[tuple[int | None, str]]:
        suffix = PurePosixPath(filename.lower()).suffix
        try:
            if suffix == ".pdf":
                reader = PdfReader(io.BytesIO(data))
                if reader.is_encrypted:
                    raise AppError("PDF có mật khẩu. Hãy tải bản không mã hóa.")
                if len(reader.pages) > 200:
                    raise AppError("Bản demo nhận tối đa 200 trang mỗi PDF.")
                pages = []
                total = 0
                for number, page in enumerate(reader.pages, 1):
                    text = (page.extract_text() or "").replace("\x00", "").strip()
                    total += len(text)
                    if total > max_chars:
                        raise AppError(f"Tài liệu vượt giới hạn {max_chars:,} ký tự.")
                    if text:
                        pages.append((number, text))
            elif suffix == ".docx":
                with ZipFile(io.BytesIO(data)) as archive:
                    if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                        raise AppError("DOCX quá lớn sau khi giải nén.")
                document = DocxDocument(io.BytesIO(data))
                parts = [p.text for p in document.paragraphs]
                parts += [
                    " | ".join(c.text for c in row.cells)
                    for table in document.tables
                    for row in table.rows
                ]
                pages = [(None, "\n".join(parts))]
            elif suffix in (".txt", ".md"):
                pages = [(None, data.decode("utf-8-sig"))]
            else:
                raise AppError("Chỉ nhận PDF có chữ, DOCX, TXT hoặc MD.", 415)
        except AppError:
            raise
        except Exception as exc:
            raise AppError("Không đọc được file. Kiểm tra định dạng hoặc dùng TXT UTF-8.") from exc
        pages = [(page, text.replace("\x00", "").strip()) for page, text in pages if text.strip()]
        if not pages:
            raise AppError("Không tìm thấy chữ. PDF scan cần OCR trước khi import.")
        if sum(len(text) for _, text in pages) > max_chars:
            raise AppError(f"Tài liệu vượt giới hạn {max_chars:,} ký tự.")
        return pages

    def chunks(self, pages: list[tuple[int | None, str]]) -> list[dict]:
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=1400,
            chunk_overlap=200,
            separators=["\n\n", "\n", ". ", " ", ""],
        )
        return [
            {"page": page, "text": part}
            for page, text in pages
            for part in splitter.split_text(text)
            if part.strip()
        ]


class IngestionService:
    def __init__(self, database, gateway, settings: Settings):
        self.db, self.gateway, self.settings = database, gateway, settings
        self.parser = DocumentParser()

    async def ingest(self, course_id, filename: str, data: bytes) -> dict:
        filename = PurePosixPath(filename.replace("\\", "/")).name[:255] or "document.txt"
        digest = hashlib.sha256(data).hexdigest()
        async with self.db.sessions() as db:
            if not await db.get(Course, course_id):
                raise AppError("Không tìm thấy môn học.", 404)
            existing = await db.scalar(
                select(Document).where(
                    Document.course_id == course_id,
                    Document.sha256 == digest,
                    Document.embedding_model == self.settings.gemini_embedding_model,
                )
            )
            if existing:
                return document_dict(existing) | {"duplicate": True}
            count = await db.scalar(
                select(func.count()).select_from(Document).where(Document.course_id == course_id)
            )
            if count >= 20:
                raise AppError("Bản demo giới hạn 20 tài liệu mỗi môn. Hãy tạo môn mới.")
        pages = await run_in_threadpool(
            self.parser.parse, filename, data, self.settings.max_document_chars
        )
        chunks = self.parser.chunks(pages)
        if len(chunks) > 250:
            raise AppError("Tài liệu có quá nhiều đoạn. Hãy chia thành các file nhỏ hơn.")
        vectors = []
        for offset in range(0, len(chunks), 8):
            vectors.extend(
                await self.gateway.embed([c["text"] for c in chunks[offset : offset + 8]])
            )
        # Publish all chunks in one transaction: failures never create a partially indexed document.
        async with self.db.sessions() as db:
            document = Document(
                course_id=course_id,
                filename=filename,
                sha256=digest,
                embedding_model=self.settings.gemini_embedding_model,
                chunk_count=len(chunks),
            )
            db.add(document)
            try:
                await db.flush()
                db.add_all(
                    [
                        Chunk(
                            document_id=document.id,
                            course_id=course_id,
                            position=i,
                            page=chunk["page"],
                            text=chunk["text"],
                            embedding=vector,
                        )
                        for i, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True))
                    ]
                )
                await db.commit()
                return document_dict(document) | {"duplicate": False}
            except IntegrityError:
                await db.rollback()
                existing = await db.scalar(
                    select(Document).where(
                        Document.course_id == course_id,
                        Document.sha256 == digest,
                        Document.embedding_model == self.settings.gemini_embedding_model,
                    )
                )
                if existing:
                    return document_dict(existing) | {"duplicate": True}
                raise


def document_dict(document: Document) -> dict:
    return {
        "id": str(document.id),
        "filename": document.filename,
        "chunk_count": document.chunk_count,
        "embedding_model": document.embedding_model,
        "created_at": document.created_at.isoformat(),
    }
